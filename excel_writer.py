"""Write parsed Mondiale FCL rates into the "FCL Rates" worksheet."""

from __future__ import annotations

import datetime as dt
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from mondiale_parser import Region

SHEET_NAME = "FCL Rates"

# (header, source field, kind) in worksheet column order A..W.
# `kind` drives cell formatting: text, number, date or note.
LAYOUT: list[tuple[str, str, str]] = [
    ("Country", "country", "text"),
    ("Origin", "origin", "text"),
    ("Destination", "destination", "text"),
    ("Ship Line", "ship_line", "text"),
    ("Service Details", "service_details", "text"),
    ("Valid From", "valid_from", "date"),
    ("Valid To", "valid_to", "date"),
    ("Notes", "notes", "note"),
    ("Currency", "currency", "text"),
    ("20'GP", "gp20", "number"),
    ("40'GP", "gp40", "number"),
    ("40'HC", "hc40", "number"),
    ("20'NOR", "nor20", "number"),
    ("40'NOR", "nor40", "number"),
    ("20'BAF", "baf20", "number"),
    ("40'BAF", "baf40", "number"),
    ("20'LSS", "lss20", "number"),
    ("40'LSS", "lss40", "number"),
    ("20'PSS/EBS", "pss20", "number"),
    ("40'PSS/EBS", "pss40", "number"),
    ("Security Fee", "sec_fee", "number"),  # header gets the currency appended
    ("20'PSC (NZD)", "psc20", "number"),
    ("40'PSC (NZD)", "psc40", "number"),
]

SECURITY_FEE_INDEX = next(i for i, (_, f, _) in enumerate(LAYOUT) if f == "sec_fee")

COLUMN_WIDTHS = {
    "Country": 18, "Origin": 20, "Destination": 18, "Ship Line": 12,
    "Service Details": 38, "Valid From": 12, "Valid To": 12, "Notes": 60,
    "Currency": 10,
}
DEFAULT_WIDTH = 11

DATE_FORMAT = "DD-MMM-YY"
NUMBER_FORMAT = "#,##0.00"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(size=10)
THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _to_number(value: str) -> float | int | str:
    """Turn a rate cell into a real number, keeping anything unparseable as text."""
    stripped = re.sub(r"[,\s$]", "", value)
    if not re.fullmatch(r"-?\d+(\.\d+)?", stripped):
        return value
    number = float(stripped)
    return int(number) if number.is_integer() else number


def _to_date(value: str) -> dt.date | str:
    try:
        return dt.datetime.strptime(value, "%d-%b-%y").date()
    except ValueError:
        return value


def _security_fee_currency(regions: list[Region]) -> str:
    """The currency the Security Fee column is quoted in, for the header label."""
    seen: dict[str, int] = {}
    for region in regions:
        for row in region.rows:
            code = row.get("sec_fee_currency", "")
            if code:
                seen[code] = seen.get(code, 0) + 1
    if not seen:
        return ""
    return max(seen, key=seen.get)


def build_workbook(regions: list[Region]) -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME

    headers = [h for h, _, _ in LAYOUT]
    currency = _security_fee_currency(regions)
    if currency:
        headers[SECURITY_FEE_INDEX] = f"Security Fee_{currency}"

    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for region in regions:
        note = region.note_text
        for row in region.rows:
            values = []
            for _, field, kind in LAYOUT:
                raw = note if field == "notes" else row.get(field, "")
                if not raw:
                    values.append(None)
                elif kind == "number":
                    values.append(_to_number(raw))
                elif kind == "date":
                    values.append(_to_date(raw))
                else:
                    values.append(raw)
            sheet.append(values)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top")
            header = headers[cell.column - 1]
            kind = LAYOUT[cell.column - 1][2]
            if kind == "date":
                cell.number_format = DATE_FORMAT
            elif kind == "number" and isinstance(cell.value, (int, float)):
                cell.number_format = NUMBER_FORMAT
            elif kind == "note":
                cell.alignment = Alignment(vertical="top", wrap_text=False)

    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = COLUMN_WIDTHS.get(header, DEFAULT_WIDTH)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"

    stream = io.BytesIO()
    workbook.save(stream)
    stream.seek(0)
    return stream


def output_filename(today: dt.date | None = None) -> str:
    today = today or dt.date.today()
    return f"Mondiale_FCL_{today:%Y-%m-%d}.xlsx"
